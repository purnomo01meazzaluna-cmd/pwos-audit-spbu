import streamlit as st
import pandas as pd
import os
from io import BytesIO
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

# Konfigurasi halaman
st.set_page_config(
    page_title="One Look Pertamina",
    page_icon="⛽",
    layout="wide"
)

# Buat folder penyimpanan bukti jika belum ada
UPLOAD_DIR = "uploaded_evidence"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# CSS Kustom untuk kartu menu & tata letak
st.markdown("""
    <style>
    .menu-card {
        background-color: white;
        border-radius: 15px;
        padding: 20px 10px;
        text-align: center;
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
        margin-bottom: 20px;
        border: 1px solid #f0f0f0;
        transition: transform 0.2s;
        display: flex;
        flex-direction: column;
        align-items: center;
        justify-content: center;
        height: 200px;
    }
    .menu-card:hover {
        transform: translateY(-5px);
        box-shadow: 0 6px 12px rgba(0, 0, 0, 0.15);
    }
    .logo-icon {
        font-size: 50px;
        margin-bottom: 10px;
    }
    .menu-title {
        font-weight: bold;
        color: #333333;
        margin-top: 5px;
        font-size: 15px;
        line-height: 1.2;
    }
    </style>
""", unsafe_allow_html=True)

# Inisialisasi state halaman & penyimpanan data audit
if 'current_page' not in st.session_state:
    st.session_state.current_page = 'dashboard'

if 'audit_data' not in st.session_state:
    st.session_state.audit_data = {}

if 'spbu_profile_data' not in st.session_state:
    st.session_state.spbu_profile_data = {
        "Nomor SPBU": "",
        "Tipe SPBU": "",
        "Alamat": "",
        "Kota/Kabupaten": "",
        "Manager/Pengelola": "",
        "No. Telepon": ""
    }

if 'kpd_data' not in st.session_state:
    st.session_state.kpd_data = {
        "mor": "",
        "sales_area": "",
        "sbm": "",
        "tipe_spbu": "",
        "tipe_kepemilikan": "",
        "tipe_audit": "",
        "asesor_1": "",
        "asesor_2": "",
        "quarter": "",
        "year": "",
        "catatan_staff": "",
        "catatan_excellence": "",
        "catatan_reliability": "",
        "catatan_visual": "",
        "catatan_extensiveness": "",
        "komentar_manager": ""
    }

# Kamus Nilai Huruf Bersama
mapping_nilai = {
    "A": 1.0,
    "B": 0.8,
    "C": 0.6,
    "D": 0.4,
    "E": 0.2,
    "F": 0.0,
    "X": 1.0
}

# Fungsi helper menentukan kategori status persentase pencapaian
def get_compliance_category(pct):
    if pct <= 35:
        return "WARNING"
    elif pct <= 60:
        return "POOR"
    elif pct <= 80:
        return "AVERAGE"
    elif pct <= 95:
        return "GOOD"
    else:
        return "EXCELLENT"

# Fungsi render item audit standar dengan Persistence (Data tersimpan & dapat diedit)
def render_audit_item(kode, pertanyaan, valid_opsi, alert_label, bobot, key_suffix, elemen_group):
    alert_badge = f" ⚠️ **[Alert: {alert_label}]**" if alert_label else ""
    st.markdown(f"**{kode}** {pertanyaan} {alert_badge}")
    st.caption(f"Valid Opsi: {valid_opsi} | Bobot: {bobot}")
    
    v_upper = valid_opsi.replace(" ", "").upper()
    if v_upper == "A/F":
        options = ["A", "F"]
    elif v_upper == "A-F":
        options = ["A", "B", "C", "D", "E", "F"]
    elif v_upper == "A/B/C/F":
        options = ["A", "B", "C", "F"]
    elif v_upper in ["A/F/X", "A/FX"]:
        options = ["A", "F", "X"]
    elif v_upper == "A/C/F":
        options = ["A", "C", "F"]
    elif v_upper == "A/C/F/X":
        options = ["A", "C", "F", "X"]
    elif v_upper == "A-F/X":
        options = ["A", "B", "C", "D", "E", "F", "X"]
    elif v_upper == "A/C/X":
        options = ["A", "C", "X"]
    elif v_upper == "A/B/F":
        options = ["A", "B", "F"]
    elif v_upper == "A/B/C/D/E/F/X":
        options = ["A", "B", "C", "D", "E", "F", "X"]
    elif v_upper == "A/B/C/D/E/F":
        options = ["A", "B", "C", "D", "E", "F"]
    else:
        options = ["A", "B", "C", "D", "E", "F", "X"]

    # Ambil nilai sebelumnya jika sudah pernah diisi
    saved_item = st.session_state.audit_data.get(kode, {})
    prev_opsi = saved_item.get("Opsi Dipilih", options[0])
    prev_catatan = saved_item.get("Catatan Assessor", "")
    if prev_catatan == "-":
        prev_catatan = ""
    
    default_index = options.index(prev_opsi) if prev_opsi in options else 0

    col_q, col_score = st.columns([5, 1])
    with col_q:
        opsi = st.radio(f"Pilih Opsi Penilaian ({kode}):", options, index=default_index, horizontal=True, key=f"opsi_{key_suffix}")
        
    nilai_huruf = mapping_nilai.get(opsi, 0.0)
    final_score = nilai_huruf * bobot
    
    with col_score:
        st.markdown(f"**Score:** `{final_score:.2f}`")
        script_alert = ""
        if alert_label and nilai_huruf < 1.0:
            st.warning(f"🚨 **Temuan Alert [{alert_label}]:** Assessor wajib melampirkan catatan atau bukti temuan.")

    col_cat, col_up = st.columns(2)
    with col_cat:
        catatan = st.text_input(f"Catatan Assessor ({kode})", value=prev_catatan, placeholder="Tambah catatan temuan...", key=f"cat_{key_suffix}")
    
    file_name_saved = saved_item.get("Lampiran Bukti Foto/Video", "Tidak Ada Lampiran")
    file_path_saved = saved_item.get("File_Path", None)
    
    with col_up:
        uploaded_file = st.file_uploader(f"Bukti Foto/Video ({kode})", type=["png", "jpg", "jpeg", "mp4", "mov"], key=f"up_{key_suffix}")
        
        if uploaded_file is not None:
            safe_kode_name = kode.replace(".", "_")
            file_extension = uploaded_file.name.split('.')[-1].lower()
            file_name_saved = f"bukti_{safe_kode_name}.{file_extension}"
            file_path_saved = os.path.join(UPLOAD_DIR, file_name_saved)
            
            with open(file_path_saved, "wb") as f:
                f.write(uploaded_file.getbuffer())
        
        # Tampilkan pratinjau jika file sudah ada di session/tersimpan
        if file_path_saved and os.path.exists(file_path_saved):
            ext = file_path_saved.split('.')[-1].lower()
            if ext in ["png", "jpg", "jpeg"]:
                st.image(file_path_saved, caption=f"Eviden Tersimpan: {kode}", use_container_width=True)
    
    st.session_state.audit_data[kode] = {
        "Code": kode,
        "Elemen": elemen_group,
        "Pertanyaan": pertanyaan,
        "Opsi Dipilih": opsi,
        "Nilai Huruf": nilai_huruf,
        "Bobot": bobot,
        "Score Total": final_score,
        "Alert": alert_label if alert_label else "-",
        "Catatan Assessor": catatan if catatan else "-",
        "Lampiran Bukti Foto/Video": file_name_saved,
        "File_Path": file_path_saved if (file_path_saved and file_path_saved.lower().endswith(('png', 'jpg', 'jpeg'))) else None
    }
    
    st.success(f"Opsi Terpilih: **{opsi}** (Nilai: {nilai_huruf} × Bobot: {bobot} = **Score: {final_score:.2f}**)")
    st.markdown("---")

# Fungsi khusus render item Densitas dengan Persistence
def render_density_audit_item(kode, nama_bbm, bobot, key_suffix):
    st.markdown(f"**{kode} Berat Jenis (densitas) {nama_bbm} diukur selama audit dalam rentang +/-0.03 dengan merujuk pada densitas dari penerimaan terakhir, yang diambil minimal 2 jam setelah pembongkaran**")
    
    saved_item = st.session_state.audit_data.get(kode, {})
    prev_opsi = saved_item.get("Opsi Dipilih", "A")
    prev_catatan = saved_item.get("Catatan Assessor", "")
    if prev_catatan == "-":
        prev_catatan = ""
        
    options_dens = ["A", "F", "X"]
    default_idx = options_dens.index(prev_opsi) if prev_opsi in options_dens else 0

    col_opt, col_cat = st.columns([2, 4])
    with col_opt:
        opsi = st.selectbox(f"Opsi ({kode})", options_dens, index=default_idx, key=f"opsi_{key_suffix}")
    with col_cat:
        catatan = st.text_input(f"Catatan / temuan untuk {kode}...", value=prev_catatan, key=f"cat_{key_suffix}")
    
    st.markdown("**Parameter Densitas (Input Asli):**")
    d1, d2, d3, d4, d5, d6 = st.columns(6)
    with d1:
        tank_no = st.text_input("Tank No.", key=f"tank_{key_suffix}")
    with d2:
        density_acuan = st.text_input("Density@15 (Acuan)", key=f"dacuan_{key_suffix}")
    with d3:
        obs_temp = st.text_input("Obs. Temp.", key=f"temp_{key_suffix}")
    with d4:
        obs_density = st.text_input("Hasil Obs. Density", key=f"odens_{key_suffix}")
    with d5:
        aktual_lapangan = st.text_input("Aktual di Lapangan", key=f"akt_{key_suffix}")
    with d6:
        st.markdown("<br>", unsafe_allow_html=True)
        st.metric(label="Density Var.", value="0.0000")

    file_name_saved = saved_item.get("Lampiran Bukti Foto/Video", "Tidak Ada Lampiran")
    file_path_saved = saved_item.get("File_Path", None)

    uploaded_file = st.file_uploader(f"Unggah Bukti ({kode})", type=["png", "jpg", "jpeg", "mp4", "mov"], key=f"up_{key_suffix}")
    if uploaded_file is not None:
        safe_kode_name = kode.replace(".", "_")
        file_extension = uploaded_file.name.split('.')[-1].lower()
        file_name_saved = f"bukti_{safe_kode_name}.{file_extension}"
        file_path_saved = os.path.join(UPLOAD_DIR, file_name_saved)
        with open(file_path_saved, "wb") as f:
            f.write(uploaded_file.getbuffer())

    nilai_huruf = mapping_nilai.get(opsi, 0.0)
    final_score = nilai_huruf * bobot

    st.session_state.audit_data[kode] = {
        "Code": kode,
        "Elemen": "Elemen 2",
        "Pertanyaan": f"Berat Jenis (densitas) {nama_bbm} diukur selama audit dalam rentang +/-0.03...",
        "Opsi Dipilih": opsi,
        "Nilai Huruf": nilai_huruf,
        "Bobot": bobot,
        "Score Total": final_score,
        "Alert": "-",
        "Catatan Assessor": catatan if catatan else "-",
        "Lampiran Bukti Foto/Video": file_name_saved,
        "File_Path": file_path_saved if (file_path_saved and file_path_saved.lower().endswith(('png', 'jpg', 'jpeg'))) else None
    }
    st.success(f"Opsi Terpilih: **{opsi}** (Score: **{final_score:.2f}**)")
    st.markdown("---")

# Fungsi khusus render item Uji Petik Nozzle (2.2.m) dengan baris dinamis
def render_nozzle_audit_item(kode, pertanyaan, bobot, key_suffix):
    st.markdown(f"**{kode} {pertanyaan}**")
    
    state_key_count = f"count_{key_suffix}"
    if state_key_count not in st.session_state:
        st.session_state[state_key_count] = 1

    saved_item = st.session_state.audit_data.get(kode, {})
    prev_opsi = saved_item.get("Opsi Dipilih", "A")
    prev_catatan = saved_item.get("Catatan Assessor", "")
    if prev_catatan == "-":
        prev_catatan = ""

    options_noz = ["A", "B", "C", "F"]
    default_idx = options_noz.index(prev_opsi) if prev_opsi in options_noz else 0

    col_opt, col_cat = st.columns([2, 4])
    with col_opt:
        opsi = st.selectbox(f"Opsi Penilaian ({kode})", options_noz, index=default_idx, key=f"opsi_{key_suffix}")
    with col_cat:
        catatan = st.text_input(f"Catatan / temuan untuk {kode}...", value=prev_catatan, key=f"cat_{key_suffix}")

    st.markdown("**Parameter Nozzle & Volume (Dapat Menambahkan Banyak Baris Nozzle):**")
    
    b_col1, b_col2, _ = st.columns([1.5, 1.5, 4])
    with b_col1:
        if st.button("➕ Tambah Baris Nozzle", key=f"add_{key_suffix}"):
            st.session_state[state_key_count] += 1
            st.rerun()
    with b_col2:
        if st.button("➖ Kurangi Baris Terakhir", key=f"del_{key_suffix}"):
            if st.session_state[state_key_count] > 1:
                st.session_state[state_key_count] -= 1
                st.rerun()

    nozzle_data_list = []
    
    for i in range(st.session_state[state_key_count]):
        st.markdown(f"**Nozzle #{i+1}**")
        n1, n2, n3, n4, n5, n6 = st.columns(6)
        with n1:
            noz_no = st.text_input("Nozzle No.", key=f"noz_no_{key_suffix}_{i}")
        with n2:
            du_make = st.text_input("DU Make", key=f"du_make_{key_suffix}_{i}")
        with n3:
            du_serial = st.text_input("DU Serial No.", key=f"du_serial_{key_suffix}_{i}")
        with n4:
            product = st.text_input("Product", key=f"prod_{key_suffix}_{i}")
        with n5:
            preset_manual = st.text_input("Preset/Manual", key=f"pres_{key_suffix}_{i}")
        with n6:
            var_ml = st.text_input("Var. (ml)", key=f"var_{key_suffix}_{i}")
        
        nozzle_data_list.append({
            "Nozzle No.": noz_no,
            "DU Make": du_make,
            "DU Serial No.": du_serial,
            "Product": product,
            "Preset/Manual": preset_manual,
            "Var. (ml)": var_ml
        })

    file_name_saved = saved_item.get("Lampiran Bukti Foto/Video", "Tidak Ada Lampiran")
    file_path_saved = saved_item.get("File_Path", None)

    uploaded_file = st.file_uploader(f"Unggah Bukti Nozzle ({kode})", type=["png", "jpg", "jpeg", "mp4", "mov"], key=f"up_{key_suffix}")
    if uploaded_file is not None:
        file_extension = uploaded_file.name.split('.')[-1].lower()
        file_name_saved = f"bukti_{kode.replace('.', '_')}.{file_extension}"
        file_path_saved = os.path.join(UPLOAD_DIR, file_name_saved)
        with open(file_path_saved, "wb") as f:
            f.write(uploaded_file.getbuffer())
        
        if file_extension in ["png", "jpg", "jpeg"]:
            st.image(uploaded_file, caption=f"Pratinjau Bukti: {kode}", use_container_width=True)

    nilai_huruf = mapping_nilai.get(opsi, 0.0)
    final_score = nilai_huruf * bobot

    st.session_state.audit_data[kode] = {
        "Code": kode,
        "Elemen": "Elemen 2",
        "Pertanyaan": pertanyaan,
        "Opsi Dipilih": opsi,
        "Nilai Huruf": nilai_huruf,
        "Bobot": bobot,
        "Score Total": final_score,
        "Alert": "UJI PETIK",
        "Catatan Assessor": catatan if catatan else "-",
        "Lampiran Bukti Foto/Video": file_name_saved,
        "File_Path": file_path_saved if (file_path_saved and file_path_saved.lower().endswith(('png', 'jpg', 'jpeg'))) else None,
        "Detail_Nozzles": nozzle_data_list
    }
    
    st.success(f"Opsi Terpilih: **{opsi}** (Score: **{final_score:.2f}**)")
    st.markdown("---")


# --- HALAMAN 1: DASHBOARD UTAMA ---
if st.session_state.current_page == 'dashboard':
    st.markdown("### ONE LOOK PERTAMINA")
    st.write("Silakan pilih kelas SPBU dan menu elemen audit di bawah ini:")
    
    kelas_spbu = st.selectbox("Pilih Kategori Kelas SPBU", ["Good", "Excellent"])
    st.write("")

    menus = [
        {"icon": "👷", "title": "Elemen 1: Skilled Staff & Services (30)", "key": "elemen1"},
        {"icon": "⛽", "title": "Elemen 2: Exact Quality & Quantity (30)", "key": "elemen2"},
        {"icon": "📋", "title": "Elemen 3: Reliable Facilities & Safety (20)", "key": "elemen3"},
        {"icon": "💬", "title": "Elemen 4: Visual Format Consistency (10)", "key": "elemen4"},
        {"icon": "📊", "title": "Elemen 5: Expansive Product Offer (10)", "key": "elemen5"},
        {"icon": "🛡️", "title": "SPBU Profile", "key": "spbu_profile"},
        {"icon": "💳", "title": "KPD (Report Metadata & Comments)", "key": "kpd"},
        {"icon": "❓", "title": "FAQ", "key": "faq"}
    ]

    rows = [menus[i:i + 3] for i in range(0, len(menus), 3)]

    for row in rows:
        cols = st.columns(3)
        for idx, menu in enumerate(row):
            with cols[idx]:
                st.markdown(f"""
                    <div class="menu-card">
                        <div class="logo-icon">{menu['icon']}</div>
                        <div class="menu-title">{menu['title']}</div>
                    </div>
                """, unsafe_allow_html=True)
                
                if st.button(f"Pilih Menu", key=menu['key'], use_container_width=True):
                    if menu['key'] == 'elemen1':
                        st.session_state.current_page = 'detail_elemen1'
                        st.rerun()
                    elif menu['key'] == 'elemen2':
                        st.session_state.current_page = 'detail_elemen2'
                        st.rerun()
                    elif menu['key'] == 'elemen3':
                        st.session_state.current_page = 'detail_elemen3'
                        st.rerun()
                    elif menu['key'] == 'elemen4':
                        st.session_state.current_page = 'detail_elemen4'
                        st.rerun()
                    elif menu['key'] == 'elemen5':
                        st.session_state.current_page = 'detail_elemen5'
                        st.rerun()
                    elif menu['key'] == 'spbu_profile':
                        st.session_state.current_page = 'detail_spbu_profile'
                        st.rerun()
                    elif menu['key'] == 'kpd':
                        st.session_state.current_page = 'detail_kpd'
                        st.rerun()
                    else:
                        st.info(f"Menu **{menu['title']}** sedang dalam pengembangan.")

    if st.session_state.audit_data:
        st.markdown("---")
        st.markdown(f"### 📊 Ringkasan Penilaian Audit SPBU (Kelas: {kelas_spbu})")
        
        score_e1 = sum([item["Score Total"] for item in st.session_state.audit_data.values() if item["Elemen"] == "Elemen 1"])
        score_e2 = sum([item["Score Total"] for item in st.session_state.audit_data.values() if item["Elemen"] == "Elemen 2"])
        score_e3 = sum([item["Score Total"] for item in st.session_state.audit_data.values() if item["Elemen"] == "Elemen 3"])
        score_e4 = sum([item["Score Total"] for item in st.session_state.audit_data.values() if item["Elemen"] == "Elemen 4"])
        score_e5 = sum([item["Score Total"] for item in st.session_state.audit_data.values() if item["Elemen"] == "Elemen 5"])
        
        total_score_all = score_e1 + score_e2 + score_e3 + score_e4 + score_e5
        total_pct = (total_score_all / 100.0) * 100

        min_total_score = 75.0 if kelas_spbu == "Good" else 80.0
        status_sertifikasi = "CERTIFIED" if total_pct >= min_total_score else "NOT CERTIFIED"

        if status_sertifikasi == "CERTIFIED":
            st.success("### STATUS: CERTIFIED ✅")
        else:
            st.error("### STATUS: NOT CERTIFIED ❌")

        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.metric(label="Total Skor Keseluruhan (TS)", value=f"{total_score_all:.2f} / 100.0 ({total_pct:.2f}%)")
        with col_m2:
            st.metric(label="Min. Nilai Kelulusan Pasti Pas", value=f"{min_total_score}%")

        st.markdown("#### Level Konfirmasi Pertamina Way (Elemen Utama)")
        
        summary_table_data = [
            {
                "Elemen Utama": "STAFF (Skilled Staff & Services)",
                "Bobot Nilai": 30.0,
                "Nilai Minimum": "85.00%",
                "Pencapaian Skor": f"{score_e1:.2f}",
                "Persentase": f"{(score_e1/30.0)*100:.2f}%" if 30.0 > 0 else "0.00%",
                "Compliance Level": get_compliance_category((score_e1/30.0)*100 if 30.0 > 0 else 0)
            },
            {
                "Elemen Utama": "EXCELLENCE (Exact Quality & Quantity)",
                "Bobot Nilai": 30.0,
                "Nilai Minimum": "85.00%",
                "Pencapaian Skor": f"{score_e2:.2f}",
                "Persentase": f"{(score_e2/30.0)*100:.2f}%" if 30.0 > 0 else "0.00%",
                "Compliance Level": get_compliance_category((score_e2/30.0)*100 if 30.0 > 0 else 0)
            },
            {
                "Elemen Utama": "RELIABILITY (Reliable Facilities & Safety)",
                "Bobot Nilai": 20.0,
                "Nilai Minimum": "85.00%",
                "Pencapaian Skor": f"{score_e3:.2f}",
                "Persentase": f"{(score_e3/20.0)*100:.2f}%" if 20.0 > 0 else "0.00%",
                "Compliance Level": get_compliance_category((score_e3/20.0)*100 if 20.0 > 0 else 0)
            },
            {
                "Elemen Utama": "VISUAL (Visual Format Consistency)",
                "Bobot Nilai": 10.0,
                "Nilai Minimum": "20.00%",
                "Pencapaian Skor": f"{score_e4:.2f}",
                "Persentase": f"{(score_e4/10.0)*100:.2f}%" if 10.0 > 0 else "0.00%",
                "Compliance Level": get_compliance_category((score_e4/10.0)*100 if 10.0 > 0 else 0)
            },
            {
                "Elemen Utama": "EXTENSIVENESS (Expansive Product Offer)",
                "Bobot Nilai": 10.0,
                "Nilai Minimum": "50.00%",
                "Pencapaian Skor": f"{score_e5:.2f}",
                "Persentase": f"{(score_e5/10.0)*100:.2f}%" if 10.0 > 0 else "0.00%",
                "Compliance Level": get_compliance_category((score_e5/10.0)*100 if 10.0 > 0 else 0)
            }
        ]
        
        st.dataframe(pd.DataFrame(summary_table_data), use_container_width=True, hide_index=True)

        st.markdown("### 📥 Unduh Data Audit")
        if st.button("Generate & Unduh Rekap Excel (.xlsx) Beserta Foto Eviden", use_container_width=True):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Rekap Audit SPBU"
            
            headers = [
                "Code", "Elemen", "Pertanyaan", "Opsi Dipilih", "Nilai Huruf", 
                "Bobot", "Score Total", "Alert", "Catatan Assessor", "Lampiran Bukti Foto/Video"
            ]
            ws.append(headers)
            
            column_widths = {'A': 10, 'B': 18, 'C': 35, 'D': 15, 'E': 12, 'F': 10, 'G': 12, 'H': 15, 'I': 25, 'J': 30}
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
            
            row_idx = 2
            for kode, data in st.session_state.audit_data.items():
                ws.append([
                    data["Code"], data["Elemen"], data["Pertanyaan"], data["Opsi Dipilih"],
                    data["Nilai Huruf"], data["Bobot"], data["Score Total"], data["Alert"],
                    data["Catatan Assessor"], ""
                ])
                ws.row_dimensions[row_idx].height = 70
                img_path = data.get("File_Path")
                if img_path and os.path.exists(img_path):
                    try:
                        img = OpenpyxlImage(img_path)
                        img.width = 100
                        img.height = 70
                        ws.add_image(img, f"J{row_idx}")
                    except Exception:
                        ws.cell(row=row_idx, column=10, value="Gagal muat gambar")
                else:
                    ws.cell(row=row_idx, column=10, value="Tidak Ada Lampiran")
                row_idx += 1
            
            excel_buffer = BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            
            st.download_button(
                label="Klik di Sini untuk Menyimpan File Excel",
                data=excel_buffer,
                file_name="rekap_audit_pertamina_dengan_eviden.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )


# --- HALAMAN SPBU PROFILE ---
elif st.session_state.current_page == 'detail_spbu_profile':
    if st.button("⬅️ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU PROFILE")
    st.write("Masukkan informasi data profil SPBU yang sedang diaudit:")
    
    with st.form("form_spbu_profile"):
        nomor_spbu = st.text_input("Nomor SPBU", value=st.session_state.spbu_profile_data["Nomor SPBU"])
        nama_spbu = st.text_input("Nama SPBU", value=st.session_state.spbu_profile_data["Nama SPBU"])
        alamat = st.text_area("Alamat Lokasi SPBU", value=st.session_state.spbu_profile_data["Alamat"])
        kota = st.text_input("Kota / Kabupaten", value=st.session_state.spbu_profile_data["Kota/Kabupaten"])
        manager = st.text_input("Nama Manager / Pengelola SPBU", value=st.session_state.spbu_profile_data["Manager/Pengelola"])
        telp = st.text_input("Nomor Telepon / Kontak", value=st.session_state.spbu_profile_data["No. Telepon"])
        
        submitted = st.form_submit_button("Simpan Profil SPBU")
        if submitted:
            st.session_state.spbu_profile_data = {
                "Nomor SPBU": nomor_spbu,
                "Nama SPBU": nama_spbu,
                "Alamat": alamat,
                "Kota/Kabupaten": kota,
                "Manager/Pengelola": manager,
                "No. Telepon": telp
            }
            st.success("Profil SPBU berhasil disimpan!")


# --- HALAMAN KPD (REPORT METADATA & AUDITOR'S COMMENTS) ---
elif st.session_state.current_page == 'detail_kpd':
    if st.button("⬅️ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## KPD (Report Metadata & Auditor's Comments)")
    st.write("Kelola informasi administratif laporan audit serta catatan/komentar asesor per elemen sesuai format lembar laporan resmi:")
    
    with st.form("form_kpd_data"):
        st.markdown("### Informasi Administratif Laporan")
        c1, c2, c3 = st.columns(3)
        with c1:
            mor = st.text_input("MOR", value=st.session_state.kpd_data["mor"])
            tipe_spbu = st.text_input("Tipe SPBU", value=st.session_state.kpd_data["tipe_spbu"])
            asesor_1 = st.text_input("Asesor 1", value=st.session_state.kpd_data["asesor_1"])
        with c2:
            sales_area = st.text_input("Sales Area", value=st.session_state.kpd_data["sales_area"])
            tipe_kepemilikan = st.text_input("Tipe Kepemilikan", value=st.session_state.kpd_data["tipe_kepemilikan"])
            asesor_2 = st.text_input("Asesor 2", value=st.session_state.kpd_data["asesor_2"])
        with c3:
            sbm = st.text_input("SBM", value=st.session_state.kpd_data["sbm"])
            tipe_audit = st.text_input("Tipe Audit", value=st.session_state.kpd_data["tipe_audit"])
            q_c, y_c = st.columns(2)
            with q_c:
                quarter = st.text_input("Quarter", value=st.session_state.kpd_data["quarter"])
            with y_c:
                year = st.text_input("Year", value=st.session_state.kpd_data["year"])

        st.markdown("---")
        st.markdown("### Auditor's Comments (Catatan Assessor per Elemen)")
        cat_staff = st.text_area("Staf Terlatih dan Termotivasi:", value=st.session_state.kpd_data["catatan_staff"])
        cat_exc = st.text_area("Jaminan Kualitas dan Kuantitas:", value=st.session_state.kpd_data["catatan_excellence"])
        cat_rel = st.text_area("Peralatan Terpelihara dan HSSE:", value=st.session_state.kpd_data["catatan_reliability"])
        cat_vis = st.text_area("Tampilan Fisik Seragam:", value=st.session_state.kpd_data["catatan_visual"])
        cat_ext = st.text_area("Portofolio Product Lengkap:", value=st.session_state.kpd_data["catatan_extensiveness"])
        kom_mgr = st.text_area("Komentar Manajer SPBU:", value=st.session_state.kpd_data["komentar_manager"])

        submitted_kpd = st.form_submit_button("Simpan Data KPD & Komentar")
        if submitted_kpd:
            st.session_state.kpd_data = {
                "mor": mor,
                "sales_area": sales_area,
                "sbm": sbm,
                "tipe_spbu": tipe_spbu,
                "tipe_kepemilikan": tipe_kepemilikan,
                "tipe_audit": tipe_audit,
                "asesor_1": asesor_1,
                "asesor_2": asesor_2,
                "quarter": quarter,
                "year": year,
                "catatan_staff": cat_staff,
                "catatan_excellence": cat_exc,
                "catatan_reliability": cat_rel,
                "catatan_visual": cat_vis,
                "catatan_extensiveness": cat_ext,
                "komentar_manager": kom_mgr
            }
            st.success("Data KPD dan Komentar berhasil disimpan!")


# --- HALAMAN 2: DETAIL ELEMEN 1 ---
elif st.session_state.current_page == 'detail_elemen1':
    if st.button("⬅ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU AUDIT CHECK LIST")
    st.markdown("### Elemen 1: Skilled Staff & Services (Max Bobot: 30)")
    
    st.markdown("### Sub-Elemen 1.1. Kebersihan dan Penampilan (10)")
    st.markdown("#### 1.1.1 Seragam (5)")
    render_audit_item("1.1.1.a", "Seluruh operator memakai seragam sesuai standar Pertamina (rancangan serupa, dikancing, baju, celana, sepatu safety warna hitam)", "A/F", "", 1.5, "1_1_1_a", "Elemen 1")
    render_audit_item("1.1.1.b", "Nama operator dan nomor SPBU tertera dan jelas terbaca", "A/F", "", 1.5, "1_1_1_b", "Elemen 1")
    render_audit_item("1.1.1.c", "Seragam dalam keadaaan bersih dan berkondisi baik", "A-F", "", 1.0, "1_1_1_c", "Elemen 1")
    render_audit_item("1.1.1.d", "Operator tidak membawa telepon genggam (HP) di area pulau pompa selama bertugas", "A/F", "HP", 1.0, "1_1_1_d", "Elemen 1")
    
    st.markdown("#### 1.1.2 Penampilan & Pemberian Hak (5)")
    render_audit_item("1.1.2.a", "Seluruh Operator berpenampilan rapi", "A-F", "", 1.0, "1_1_2_a", "Elemen 1")
    render_audit_item("1.1.2.b", "Seluruh operator menerima upah sesuai aturan Upah Minimum dan mendapatkan benefit", "A/B/C/F", "", 2.0, "1_2_b_hak", "Elemen 1")
    render_audit_item("1.1.2.c", "Pekerja mendapat bagian sesuai haknya dari Program Reward PT Pertamina Patra Niaga", "A/F/X", "", 2.0, "1_2_c_rew", "Elemen 1")

    st.markdown("### Sub-Elemen 1.2 Prosedur Pelayanan (20)")
    render_audit_item("1.2.a", "Pelanggan disambut dengan sopan serta salam, ditawarkan produk JBU Top Tier", "A/C/F", "Salam", 3.0, "1_2_a", "Elemen 1")
    render_audit_item("1.2.b", "Operator mengingatkan & memastikan mesin kendaraan konsumen dalam keadaan mati", "A/C/F", "", 2.5, "1_2_b", "Elemen 1")
    render_audit_item("1.2.c", "Penunjuk angka meter dimulai dari angka 'nol'", "A/C/F/X", "Nol", 3.0, "1_2_c", "Elemen 1")
    render_audit_item("1.2.d", "Pengisian BBM dilakukan secara hati-hati untuk mencegah tumpah", "A/F/X", "", 2.0, "1_2_d", "Elemen 1")
    render_audit_item("1.2.e", "Operator menawarkan pembayaran menggunakan aplikasi MyPertamina", "A/F", "", 2.5, "1_2_e", "Elemen 1")
    render_audit_item("1.2.f", "Operator mengkonfirmasi harga total dan jumlah uang", "A-F", "", 1.5, "1_2_f", "Elemen 1")
    render_audit_item("1.2.g", "Operator menyerahkan kuitansi/Struk dan kembalian", "A-F/X", "", 1.5, "1_2_g", "Elemen 1")
    render_audit_item("1.2.h", "Operator mengucapkan terima kasih kepada pelanggan", "A/C/F", "Trims", 3.0, "1_2_h", "Elemen 1")
    render_audit_item("1.2.i", "Tersedia informasi Call Center Layanan Pelanggan", "A/F", "", 1.0, "1_2_i", "Elemen 1")


# --- HALAMAN 3: DETAIL ELEMEN 2 ---
elif st.session_state.current_page == 'detail_elemen2':
    if st.button("⬅️ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU Q & Q CHECK")
    st.markdown("### Elemen 2: Exact Quality & Quantity (Max Bobot: 30)")
    
    st.markdown("### Sub-Elemen 2.1: Peralatan (7)")
    render_audit_item("2.1.a", "Dispenser Unit disegel dan disertifikasi oleh Dinas Metrologi", "A/F", "", 2.5, "2_1_a", "Elemen 2")
    render_audit_item("2.1.b", "SPBU memperbaharui secara berkala catatan Totalizer Dispenser Unit BBM", "A/F", "", 2.0, "2_1_b", "Elemen 2")
    render_audit_item("2.1.c", "Seluruh peralatan Q&Q tersedia dan dalam kondisi baik", "A/F", "", 2.5, "2_1_c", "Elemen 2")

    st.markdown("### Sub-Elemen 2.2 Prosedur Monitoring (23)")
    render_audit_item("2.2.a", "Tidak ditemukan tanda-tanda manipulasi pada dispenser unit", "A/F", "DU Dispenser", 1.0, "2_2_a", "Elemen 2")
    render_audit_item("2.2.b", "Sampel 2 pengiriman terakhir dari tiap jenis BBM disimpan dalam kontainer aluminium", "A/F", "", 0.75, "2_2_b", "Elemen 2")
    render_audit_item("2.2.c", "Kaleng Sampel disegel dan Label Sampel BBM sudah ditempel", "A/F", "", 0.5, "2_2_c", "Elemen 2")
    render_audit_item("2.2.d", "Tersedia Display sampel BBM yang sesuai standar Pertamina", "A/F", "", 0.25, "2_2_d", "Elemen 2")
    render_audit_item("2.2.e", "Tidak ditemukan air dalam tangki-tangki timbun BBM", "A/F", "", 1.0, "2_2_e", "Elemen 2")
    
    render_density_audit_item("2.2.f", "Pertalite (Oktan 90)", 1.0, "2_2_f")
    render_density_audit_item("2.2.g", "Pertamax (Oktan 92)", 1.0, "2_2_g")
    render_density_audit_item("2.2.h", "Pertamax Green", 1.0, "2_2_h")
    render_density_audit_item("2.2.i", "Pertamax Turbo", 1.0, "2_2_i")
    render_density_audit_item("2.2.j", "Bio Solar/Solar", 1.0, "2_2_j")
    render_density_audit_item("2.2.k", "Pertamina Dex", 1.0, "2_2_k")
    render_density_audit_item("2.2.l", "Dexlite", 1.0, "2_2_l")

    pertanyaan_22m = (
        "Volume BBM yang dikeluarkan dari nozzle yang diperiksa secara manual dan program "
        "berada dalam rentang toleransi (- 60 ml dengan bejana ukur 20 liter). "
        "(100% dari jumlah nozzle yang ada untuk SPBU Excellent dan 50% dari jumlah nozzle "
        "yang ada untuk SPBU Good, dari masing-masing produk diperiksa oleh Auditor secara acak)"
    )
    render_nozzle_audit_item("2.2.m", pertanyaan_22m, 9.5, "2_2_m")

    render_audit_item("2.2.n", "Catatan stok harian disimpan dan selalu diperbarui", "A/C/F", "", 1.0, "2_2_n", "Elemen 2")
    render_audit_item("2.2.o", "Catatan kualitas harian dan Pemeriksaan Volume Harian disimpan dan selalu diperbarui", "A/C/F", "", 1.0, "2_2_o", "Elemen 2")
    render_audit_item("2.2.p", "Tanda terima dan Penebusan BBM tersedia", "A/F", "", 0.5, "2_2_p", "Elemen 2")
    render_audit_item("2.2.q", "Semua produk JBU yang ditawarkan tersedia", "A/F", "", 0.5, "2_2_q", "Elemen 2")


# --- HALAMAN 4: DETAIL ELEMEN 3 ---
elif st.session_state.current_page == 'detail_elemen3':
    if st.button("⬅ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU AUDIT CHECK LIST")
    st.markdown("### Elemen 3: Reliable Facilities & Safety (Max Bobot: 20)")
    
    st.markdown("### Sub-Elemen 3.1: Kebersihan Harian (14.5)")
    st.markdown("#### 3.1.1 Halaman Depan (6.5)")
    render_audit_item("3.1.1.a", "Driveway/ Pelataran pengisian BBM dalam keadaan bebas tumpahan minyak, sampah, kering, dan terpelihara baik (tak ada lubang atau genangan air)", "A-F", "", 0.9, "3_1_1_a", "Elemen 3")
    render_audit_item("3.1.1.b", "Pulau pompa dan kolom-kolom kanopi dalam kondisi bebas minyak atau tanda kotor serta tak terlihat kerusakan", "A-F", "", 0.75, "3_1_1_b", "Elemen 3")
    render_audit_item("3.1.1.c", "Dispenser Unit BBM dalam kondisi bebas dari minyak, cat yang rusak atau coretan", "A-F", "", 0.75, "3_1_1_c", "Elemen 3")
    render_audit_item("3.1.1.d", "SPBU tersedia Oil Spill Kit", "A-F", "", 0.2, "3_1_1_d", "Elemen 3")
    render_audit_item("3.1.1.e", "Totem/Signboard, Lisplang, Kanopi, Rambu Masuk dan Rambu Keluar dalam kondisi bersih dan baik tanpa ada kerusakan", "A-F", "", 1.0, "3_1_1_e", "Elemen 3")
    render_audit_item("3.1.1.f", "Seluruh lampu penerangan di bawah kanopi dan lampu lisplang dalam kondisi berfungsi", "A-F", "", 0.4, "3_1_1_f", "Elemen 3")
    render_audit_item("3.1.1.g", "Seluruh lampu halaman SPBU berfungsi dengan baik (area di sekitar batas SPBU)", "A-F", "", 0.2, "3_1_1_g", "Elemen 3")
    render_audit_item("3.1.1.h", "Area penyimpanan BBM (area tangki timbun) dalam kondisi bersih, tidak tampak tumpahan BBM, dan sampah seperti sobekan kertas", "A-F", "", 0.3, "3_1_1_h", "Elemen 3")
    render_audit_item("3.1.1.i", "Pelataran/tempat pembongkaran BBM dalam kondisi baik, tidak ada lubang", "A-F", "", 0.2, "3_1_1_i", "Elemen 3")
    render_audit_item("3.1.1.j", "Tutup lubang pengisian BBM (Oil Sump) diberi kode warna sesuai produk", "A-F", "", 0.4, "3_1_1_j", "Elemen 3")
    render_audit_item("3.1.1.k", "Lubang pengisian BBM (Oil Sump) bersih dari tumpahan minyak", "A/C/F", "", 0.3, "3_1_1_k", "Elemen 3")
    render_audit_item("3.1.1.l", "Oil Catcher (Jebakan minyak) sesuai standar dan dalam kondisi bersih", "A/C/F", "", 0.5, "3_1_1_l", "Elemen 3")
    render_audit_item("3.1.1.m", "Taman dalam kondisi bersih", "A/C/F/X", "", 0.4, "3_1_1_m", "Elemen 3")
    render_audit_item("3.1.1.n", "Lampu penerangan taman berfungsi dengan baik", "A/F/X", "", 0.2, "3_1_1_n", "Elemen 3")

    st.markdown("#### 3.1.2 Toilet (4.5)")
    render_audit_item("3.1.2.a", "Toilet sesuai standar PT Pertamina Patra Niaga", "A/C", "", 0.5, "3_1_2_a", "Elemen 3")
    render_audit_item("3.1.2.b", "Toilet dan wastafel dalam kondisi bersih, terpelihara baik dan perangkat di dalamnya dapat berfungsi (seperti kran air, bak air, gayung)", "A/C/F", "", 1.25, "3_1_2_b", "Elemen 3")
    render_audit_item("3.1.2.c", "Tidak ada kerusakan yang terlihat pada perangkat utama toilet (seperti dinding, plafon, pintu, jendela, gantungan, dan keramik)", "A/F", "", 1.0, "3_1_2_c", "Elemen 3")
    render_audit_item("3.1.2.d", "Akses mudah menuju toilet disertai tanda penunjuk (signage) yang diperlukan dan berpenerangan cukup", "A/F", "", 0.5, "3_1_2_d", "Elemen 3")
    render_audit_item("3.1.2.e", "Toilet berpenerangan cukup dan tersedia pengharum ruangan", "A/F", "", 1.0, "3_1_2_e", "Elemen 3")
    render_audit_item("3.1.2.f", "Tempat sampah dan keset tersedia pada setiap ruangan toilet", "A/F", "", 0.25, "3_1_2_f", "Elemen 3")

    st.markdown("#### 3.1.3 Fasilitas Ibadah (1)")
    render_audit_item("3.1.3.a", "Fasilitas ibadah sesuai standar PT Pertamina Patra Niaga", "A/C/X", "", 0.5, "3_1_3_a", "Elemen 3")
    render_audit_item("3.1.3.b", "Area untuk wudhu dan musala dalam keadaan bersih, berpenerangan cukup, terpelihara dengan baik serta tersedia pengharum ruangan", "A/C/F/X", "", 0.2, "3_1_3_b", "Elemen 3")
    render_audit_item("3.1.3.c", "Tidak ada kerusakan yang terlihat pada perangkat utama area wudhu dan musala (seperti dinding, plafon, pintu, jendela, kran air, dan lantai)", "A/C/F/X", "", 0.15, "3_1_3_c", "Elemen 3")
    render_audit_item("3.1.3.d", "Alat perangkat salat tersedia di musala, dalam kondisi baik dan segar", "A/C/F/X", "", 0.15, "3_1_3_d", "Elemen 3")

    st.markdown("#### 3.1.4 Aspek HSSE (2.5)")
    render_audit_item("3.1.4.a", "Tersedia alat pemadam api ringan (APAR) dalam kondisi baik dan mudah diakses", "A/F", "APAR", 0.1, "3_1_4_a", "Elemen 3")
    render_audit_item("3.1.4.b", "Tersedia minimal 2 (dua) buah alat pemadam api beroda (APAP)", "A/F", "APAP", 0.1, "3_1_4_b", "Elemen 3")
    render_audit_item("3.1.4.c", "Alat pemadam api memiliki kartu inspeksi dan masih berlaku", "A/F", "Masa Berlaku APAR", 0.1, "3_1_4_c", "Elemen 3")
    render_audit_item("3.1.4.d", "\"Grounding\" di area pengisian BBM dalam kondisi baik", "A/F", "Grounding", 0.1, "3_1_4_d", "Elemen 3")
    render_audit_item("3.1.4.e", "Ducting / saluran pipa dari tangki timbun ke dispenser ditimbun", "A/F", "Ducting", 0.1, "3_1_4_e", "Elemen 3")
    render_audit_item("3.1.4.f", "Seluruh nozzle terpasang Breakaway Valve", "A/F", "Breakaway Valve", 0.1, "3_1_4_f", "Elemen 3")
    render_audit_item("3.1.4.g", "Dispensing Sump & Impact Valve terpasang dan dalam kondisi baik", "A/F", "", 0.2, "3_1_4_g", "Elemen 3")
    render_audit_item("3.1.4.h", "Tidak terdapat kegiatan lain (jualan) di zona berbahaya", "A/F", "", 0.1, "3_1_4_h", "Elemen 3")
    render_audit_item("3.1.4.i", "Terdapat stop kontak di kanopi yang sesuai spesifikasi (tertutup)", "A/F", "", 0.2, "3_1_4_i", "Elemen 3")
    render_audit_item("3.1.4.j", "Emergency shut down tersedia di setiap dispenser dan berfungsi baik", "A/F", "Emergency Shut Down", 0.1, "3_1_4_j", "Elemen 3")
    render_audit_item("3.1.4.k", "Junction box kabel kedap", "A/F", "", 0.2, "3_1_4_k", "Elemen 3")
    render_audit_item("3.1.4.l", "Tidak ada lubang terbuka di area manhole", "A/F", "ST Manhole", 0.1, "3_1_4_l", "Elemen 3")
    render_audit_item("3.1.4.m", "Tidak terdapat genangan BBM / air di dalam tank sump", "A/F", "ST Air", 0.1, "3_1_4_m", "Elemen 3")
    render_audit_item("3.1.4.n", "Tersedia rambu-rambu / sticker peringatan safety", "A/F", "", 0.1, "3_1_4_n", "Elemen 3")
    render_audit_item("3.1.4.o", "Tersedia daftar nomor telepon penting / emergency", "A/F", "", 0.1, "3_1_4_o", "Elemen 3")
    render_audit_item("3.1.4.p", "Tersedia Surat Ijin Kerja Aman (SIKA)", "A/F/X", "", 0.1, "3_1_4_p", "Elemen 3")
    render_audit_item("3.1.4.q", "SPBU tersedia dokumen UKL/UPL", "A/C", "", 0.1, "3_1_4_q", "Elemen 3")
    render_audit_item("3.1.4.r", "Instalasi listrik sesuai dengan standar", "A/F", "", 0.1, "3_1_4_r", "Elemen 3")
    render_audit_item("3.1.4.s", "Kotak P3K tersedia di kantor", "A/F", "", 0.1, "3_1_4_s", "Elemen 3")
    render_audit_item("3.1.4.t", "Operator dan pengawas telah terlatih pemadaman kebakaran", "A/F", "", 0.2, "3_1_4_t", "Elemen 3")
    render_audit_item("3.1.4.u", "Terdapat minimal 1 petugas berlisensi Safetyman", "A/F", "Safetyman", 0.1, "3_1_4_u", "Elemen 3")

    st.markdown("### Sub-Elemen 3.2: Pemeliharaan berkala atas DU, ST, dan Fasilitas (4.5)")
    render_audit_item("3.2.a", "Catatan pemeliharaan Fasilitas SPBU diperbarui", "A/F", "", 0.5, "3_2_a", "Elemen 3")
    render_audit_item("3.2.b", "Catatan pemeliharaan DU & ST diperbarui", "A/F", "", 0.5, "3_2_b", "Elemen 3")
    render_audit_item("3.2.c", "Dispenser Unit tidak tampak kerusakan/cat terkelupas", "A/F", "", 0.5, "3_2_c", "Elemen 3")
    render_audit_item("3.2.d", "Layar penunjuk (LCD Dispenser) terbaca jelas", "A/F", "", 0.5, "3_2_d", "Elemen 3")
    render_audit_item("3.2.e", "Tidak ada kebocoran pipa produk di dalam DU", "A/F", "", 0.5, "3_2_e", "Elemen 3")
    render_audit_item("3.2.f", "Koneksi listrik di dalam DU aman", "A/F", "", 0.5, "3_2_f", "Elemen 3")
    render_audit_item("3.2.g", "Selang pengisian BBM tidak bocor/terkelupas", "A/F", "", 0.5, "3_2_g", "Elemen 3")
    render_audit_item("3.2.h", "Generator terpelihara dan berfungsi baik", "A/F", "", 0.5, "3_2_h", "Elemen 3")
    render_audit_item("3.2.i", "Pipa sirkulasi udara tangki (Vent Pipe) sesuai warna produk", "A/F/X", "", 0.5, "3_2_i", "Elemen 3")

    st.markdown("### Sub-Elemen 3.3: Uraian Pemeliharaan Kerusakan (1)")
    render_audit_item("3.3.a", "Catatan pemeliharaan kerusakan tersedia dan terpelihara", "A/F", "", 0.25, "3_3_a", "Elemen 3")
    render_audit_item("3.3.b", "Tanggal keluhan kerusakan dan penanganan jelas", "A/F", "", 0.25, "3_3_b", "Elemen 3")
    render_audit_item("3.3.c", "Keluhan diselesaikan dalam waktu maksimal 3 bulan", "A/F", "", 0.25, "3_3_c", "Elemen 3")
    render_audit_item("3.3.d", "Seluruh mesin dalam kondisi baik dan berfungsi", "A/F", "", 0.25, "3_3_d", "Elemen 3")


# --- HALAMAN 5: DETAIL ELEMEN 4 ---
elif st.session_state.current_page == 'detail_elemen4':
    if st.button("⬅️ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU AUDIT CHECK LIST")
    st.markdown("### Elemen 4: Visual Format Consistency (Max Bobot: 10)")
    
    st.markdown("### Sub-Elemen 4.1: Identitas Visual Ritel (4)")
    render_audit_item("4.1.a", "Produk Sign sesuai standar PERTAMINA", "A/F", "", 1.0, "4_1_a", "Elemen 4")
    render_audit_item("4.1.b", "Totem sesuai standar PERTAMINA", "A-F", "", 1.0, "4_1_b", "Elemen 4")
    render_audit_item("4.1.c", "Listplank (Facia) sesuai standar PERTAMINA", "A/F", "", 1.0, "4_1_c", "Elemen 4")
    render_audit_item("4.1.d", "Tiang kanopi sesuai standar PERTAMINA", "A/F", "", 1.0, "4_1_d", "Elemen 4")

    st.markdown("### Sub-Elemen 4.2: Dispenser Unit (2)")
    render_audit_item("4.2.a", "Dispenser Unit memiliki warna penunjuk produk standar", "A/F", "", 1.0, "4_2_a", "Elemen 4")
    render_audit_item("4.2.b", "Paduan warna Dispenser Unit sesuai standar", "A/F", "", 1.0, "4_2_b", "Elemen 4")

    st.markdown("### Sub-Elemen 4.3: Lain-lain (4)")
    render_audit_item("4.3.a", "EDC Digitalisasi / Tablet MyPertamina tersedia dan berfungsi", "A/B/C/F", "EDC", 0.5, "4_3_a", "Elemen 4")
    render_audit_item("4.3.b", "Petunjuk fasilitas SPBU telah tersedia", "A/F", "", 0.5, "4_3_b", "Elemen 4")
    render_audit_item("4.3.c", "Penempatan Signage Tenant sesuai ketentuan", "A/F", "", 0.5, "4_3_c", "Elemen 4")
    render_audit_item("4.3.d", "SPBU menggunakan ATG & POS sesuai ketentuan", "A/F", "", 1.0, "4_3_d", "Elemen 4")
    render_audit_item("4.3.e", "CCTV di setiap pulau pompa berfungsi dan arsip min. 1 bulan", "A/F", "CCTV", 0.5, "4_3_e", "Elemen 4")
    render_audit_item("4.3.f", "Jalur Red Carpet Fast Track tersedia dan sesuai", "A/C/F", "Jalur Fast Track", 0.5, "4_3_f", "Elemen 4")
    render_audit_item("4.3.g", "Terdapat Dedicated Operator dengan Rompi Khusus", "A/F", "", 0.25, "4_3_g", "Elemen 4")
    render_audit_item("4.3.h", "Posisi jalur Red Carpet mudah diakses", "A/C/F", "", 0.25, "4_3_h", "Elemen 4")


# --- HALAMAN 6: DETAIL ELEMEN 5 ---
elif st.session_state.current_page == 'detail_elemen5':
    if st.button("⬅️ Kembali ke Menu Utama"):
        st.session_state.current_page = 'dashboard'
        st.rerun()
    
    st.markdown("---")
    st.markdown("## SPBU AUDIT CHECK LIST")
    st.markdown("### Elemen 5: Expansive Product Offer (Max Bobot: 10)")
    
    st.markdown("### Sub-Elemen 5.1: Penawaran BBM")
    render_audit_item("5.1.a", "Tersedianya produk Pertamax Turbo", "A/F", "", 0.35, "5_1_a", "Elemen 5")
    render_audit_item("5.1.b", "Tersedianya produk Pertamax Green", "A/F", "", 0.3, "5_1_b", "Elemen 5")
    render_audit_item("5.1.c", "Tersedianya produk Pertamax", "A/F", "CPO", 0.23, "5_1_c", "Elemen 5")
    render_audit_item("5.1.d", "Tersedianya produk Pertamina Dex", "A/F", "CPO", 0.3, "5_1_d", "Elemen 5")
    render_audit_item("5.1.e", "Tersedianya produk Dexlite", "A/F", "", 0.23, "5_1_e", "Elemen 5")
    render_audit_item("5.1.f", "Tersedia produk JBU minimum 1 jenis (Pertamax/Dex Series)", "A/C/F", "Produk JBU", 0.21, "5_1_f", "Elemen 5")
    render_audit_item("5.1.g", "Realisasi penebusan JBU per-2 bulan sesuai target region", "A/C/F", "VOLUME JBU", 0.2, "5_1_g", "Elemen 5")
    render_audit_item("5.1.h", "Kesesuaian materi promo dengan program berjalan", "A/F/X", "", 0.18, "5_1_h", "Elemen 5")

    st.markdown("### Sub-Elemen 5.2: Penawaran non-BBM")
    render_audit_item("5.2.a", "SPBU tersedia NFR brand Bright", "A/B/F", "", 0.8, "5_2_a", "Elemen 5")
    render_audit_item("5.2.b", "SPBU tersedia NFR Internasional", "A/B/F", "NFR INT", 0.65, "5_2_b", "Elemen 5")
    render_audit_item("5.2.c", "SPBU tersedia fasilitas EBT (PLTS / Layanan EV)", "A/B/F", "", 1.15, "5_2_c", "Elemen 5")
    render_audit_item("5.2.d", "SPBU tersedia NFR Nasional", "A/B/F", "", 1.15, "5_2_d", "Elemen 5")
    render_audit_item("5.2.e", "SPBU tersedia NFR Lokal berizin payung Pertamina", "A/B/F", "NFR LKL", 0.9, "5_2_e", "Elemen 5")
    render_audit_item("5.2.f", "Seluruh NFR memiliki Izin Prinsip", "A/F", "Izin Prinsip", 0.2, "5_2_f", "Elemen 5")
    render_audit_item("5.2.g", "Kelengkapan NFR sesuai kategori kelas SPBU", "A/B/F", "NFR", 1.55, "5_2_g", "Elemen 5")
    render_audit_item("5.2.h", "Tersedia 2 brand pelumas (Fastron & Enduro series)", "A/B/F", "PELUMAS", 0.8, "5_2_h", "Elemen 5")
    render_audit_item("5.2.i", "Tersedia fasilitas pengisian air dan angin", "A/F", "", 0.8, "5_2_i", "Elemen 5")
